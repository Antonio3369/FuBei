#!/usr/bin/env python3
"""付呗交易监控。本地打开浏览器，按导出总表跟进沉默商户、掉量和额度包。"""

from __future__ import annotations

import json
import re
import sys
import webbrowser
from datetime import date, datetime, time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
WEB = ROOT / "web"
DATA = ROOT / "data"
SNAPSHOTS = DATA / "snapshots"
UPLOADS = DATA / "uploads"
SEED = Path("/Users/xin/Desktop/付呗/碰一下音响商户交易数据统计10.11.xlsx")

def num(value) -> float:
    if isinstance(value, bool) or value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def text(value) -> str:
    if value is None or isinstance(value, (datetime, time, date)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, bool):
        return ""
    raw = str(value).strip()
    if raw in ("", "0", "0.0", "/"):
        return ""
    return raw


def opened_value(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    return None


def parse_quota(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        shown = str(int(value)) if float(value).is_integer() else str(value)
        return shown, "configured"
    raw = "" if value is None else str(value).strip()
    if raw in ("", "/", "0", "未参加银行活动"):
        return "", "none"
    if raw == "待配置":
        return "待配置", "pending"
    return raw, "configured"


def parse_bank(value):
    raw = text(value)
    if raw in ("", "未参加银行活动"):
        return "未参加银行活动", False
    return raw, True


def infer_as_of(filename: str, opened_dates: list[date]) -> date:
    year = max((item.year for item in opened_dates), default=date.today().year)
    match = re.search(r"(\d{1,2})月(\d{1,2})", filename) or re.search(r"(?<!\d)(\d{1,2})[.\-](\d{1,2})(?!\d)", filename)
    if match:
        month, day = int(match.group(1)), int(match.group(2))
        try:
            return date(year, month, day)
        except ValueError:
            pass
    if opened_dates:
        return max(opened_dates)
    return date.today()


def load_sheets(path: Path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            headers = [str(cell).strip() if cell is not None else "" for cell in header]
            if "设备SN" not in headers:
                continue
            records = []
            for row in rows:
                if not row or row[0] in (None, ""):
                    continue
                item = {}
                for index, key in enumerate(headers):
                    if key and index < len(row):
                        item[key] = row[index]
                item["设备SN"] = str(item["设备SN"]).strip()
                records.append(item)
            sheets.append({"name": name, "has_company": "所属公司" in headers, "rows": records})
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("表格里没有「设备SN」列，请使用付呗总表。")
    return sheets


def merge_sheets(sheets):
    primary = next((item for item in sheets if item["has_company"]), None)
    if primary is None:
        primary = max(sheets, key=lambda item: len(item["rows"]))
    merged = {}
    order = []
    for record in primary["rows"]:
        serial = record["设备SN"]
        if serial not in merged:
            order.append(serial)
        merged[serial] = record
    for item in sheets:
        if item["name"] == primary["name"]:
            continue
        for record in item["rows"]:
            serial = record["设备SN"]
            if serial not in merged:
                merged[serial] = record
                order.append(serial)
    return [merged[serial] for serial in order]


def normalize(records, filename: str):
    opened_dates = []
    for record in records:
        opened = opened_value(record.get("开户时间"))
        if opened:
            opened_dates.append(opened.date())
    as_of = infer_as_of(filename, opened_dates)
    rows = []
    for record in records:
        opened = opened_value(record.get("开户时间"))
        deploy = text(record.get("设备铺设方式")) or "未知"
        bank, joined = parse_bank(record.get("报名银行政策"))
        quota, quota_status = parse_quota(record.get("额度包配置情况"))
        month_count = int(num(record.get("本月交易笔数")))
        month_amount = round(num(record.get("本月交易金额")), 2)
        month_alipay = int(num(record.get("本月支付宝笔数")))
        month_tap = int(num(record.get("本月碰一下笔数")))
        week_count = int(num(record.get("上周交易笔数")))
        week_amount = round(num(record.get("上周交易金额")), 2)
        week_alipay = int(num(record.get("上周支付宝笔数")))
        week_tap = int(num(record.get("上周碰一下笔数")))
        company = text(record.get("所属公司"))
        if not company and deploy == "未铺设":
            company = "未铺设"
        manager = text(record.get("授理商")) or text(record.get("所属经理"))
        sales = text(record.get("市场经理")) or text(record.get("所属业务员"))
        activated = opened is not None
        age = (as_of - opened.date()).days if opened else None
        if age is not None and age < 0:
            age = 0
        flags = []
        if deploy == "未铺设":
            flags.append("undeployed")
        elif not activated:
            flags.append("incomplete")
        if month_count == 0 and deploy != "未铺设":
            flags.append("silent")
        if activated and month_count == 0 and age is not None and age <= 2:
            flags.append("new_silent")
        if quota_status == "pending":
            flags.append("quota_pending")
        if month_count > 0 and month_alipay == 0:
            flags.append("no_alipay")
            if joined:
                flags.append("no_alipay_bank")
        if 0 < month_count <= 3:
            flags.append("low")
        if 0 < month_count <= 3 and month_amount >= 10000:
            flags.append("large_ticket")
        rows.append(
            {
                "sn": record.get("设备SN") or "",
                "company": company,
                "manager": manager,
                "sales": sales,
                "merchant_id": text(record.get("商户ID")),
                "opened_at": opened.strftime("%Y-%m-%d %H:%M") if opened else "",
                "opened_date": opened.strftime("%Y-%m-%d") if opened else "",
                "age": age,
                "account": text(record.get("商户账号")),
                "store_id": text(record.get("门店id")),
                "store_name": text(record.get("门店名称")),
                "deploy": deploy,
                "activated": activated,
                "bank": bank,
                "joined_bank": joined,
                "quota": quota,
                "quota_status": quota_status,
                "month_count": month_count,
                "month_amount": month_amount,
                "month_alipay": month_alipay,
                "month_tap": month_tap,
                "week_count": week_count,
                "week_amount": week_amount,
                "week_alipay": week_alipay,
                "week_tap": week_tap,
                "flags": flags,
            }
        )
    return as_of, rows


def slim_previous(previous):
    if not previous:
        return None
    kept = ("silent", "quota_pending")
    rows = []
    for row in previous["rows"]:
        rows.append(
            {
                "sn": row["sn"],
                "company": row.get("company") or "",
                "manager": row.get("manager") or "",
                "sales": row.get("sales") or "",
                "month_amount": row.get("month_amount") or 0,
                "month_count": row.get("month_count") or 0,
                "week_amount": row.get("week_amount") or 0,
                "week_count": row.get("week_count") or 0,
                "activated": bool(row.get("activated")),
                "flags": [flag for flag in row.get("flags") or [] if flag in kept],
            }
        )
    return {"as_of": previous["as_of"], "rows": rows}


def snapshot_history(as_of: str):
    if not SNAPSHOTS.exists():
        return []
    files = sorted(path for path in SNAPSHOTS.glob("*.json") if path.stem < as_of)
    return [slim_previous(read_snapshot(path)) for path in files]


def with_context(payload, previous=None):
    history = snapshot_history(payload["meta"]["as_of"])
    payload["history"] = history
    latest = history[-1] if history else None
    payload["previous"] = latest
    payload["diff"] = snapshot_diff(latest, payload["rows"]) if latest else None
    return payload


def snapshot_diff(previous, current):
    if not previous:
        return None
    prev_rows = {row["sn"]: row for row in previous["rows"]}
    curr_rows = {row["sn"]: row for row in current}
    prev_amount = round(sum(row["month_amount"] for row in previous["rows"]), 2)
    curr_amount = round(sum(row["month_amount"] for row in current), 2)
    prev_count = sum(row["month_count"] for row in previous["rows"])
    curr_count = sum(row["month_count"] for row in current)

    def has(row, flag):
        return flag in row["flags"]

    newly_silent = 0
    resolved_silent = 0
    new_quota = 0
    for serial, row in curr_rows.items():
        old = prev_rows.get(serial)
        if old is None:
            continue
        if has(row, "silent") and not has(old, "silent"):
            newly_silent += 1
        if has(old, "silent") and not has(row, "silent"):
            resolved_silent += 1
        if has(row, "quota_pending") and not has(old, "quota_pending"):
            new_quota += 1
    new_activated = sum(1 for serial, row in curr_rows.items() if row["activated"] and serial not in prev_rows)
    return {
        "prev_as_of": previous["as_of"],
        "amount_delta": round(curr_amount - prev_amount, 2),
        "count_delta": curr_count - prev_count,
        "new_activated": new_activated,
        "newly_silent": newly_silent,
        "resolved_silent": resolved_silent,
        "new_quota_pending": new_quota,
    }


def build_from_xlsx(path: Path):
    sheets = load_sheets(path)
    records = merge_sheets(sheets)
    as_of, rows = normalize(records, path.name)
    sample = records[0] if records else {}
    payload = {
        "meta": {
            "source": path.name,
            "as_of": as_of.isoformat(),
            "devices": len(rows),
            "sheets": [item["name"] for item in sheets],
            "has_company": "所属公司" in sample,
            "has_tap": "本月碰一下笔数" in sample,
            "manager_label": "授理商" if "授理商" in sample else "经理",
            "sales_label": "市场经理" if "市场经理" in sample else "业务员",
            "product": "碰一下音响" if "本月碰一下笔数" in sample else "付呗设备",
        },
        "rows": rows,
    }
    return with_context(payload)


def read_snapshot(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def save_snapshot(payload):
    SNAPSHOTS.mkdir(parents=True, exist_ok=True)
    target = SNAPSHOTS / f"{payload['meta']['as_of']}.json"
    stored = {"as_of": payload["meta"]["as_of"], "source": payload["meta"]["source"], "rows": payload["rows"]}
    target.write_text(json.dumps(stored, ensure_ascii=False), encoding="utf-8")
    return target


def ensure_seed():
    DATA.mkdir(parents=True, exist_ok=True)
    existing = sorted(DATA.glob("*.xlsx"), key=lambda item: item.stat().st_mtime, reverse=True)
    if existing:
        return existing[0]
    if SEED.exists():
        target = DATA / SEED.name
        target.write_bytes(SEED.read_bytes())
        return target
    return None


def load_current():
    path = ensure_seed()
    if path is None:
        return {
            "meta": {"source": "", "as_of": "", "devices": 0, "sheets": []},
            "rows": [],
            "diff": None,
        }
    # Peek date from an existing snapshot of the same file only after parse.
    payload = build_from_xlsx(path)
    save_snapshot(payload)
    return payload


def parse_upload(body: bytes, content_type: str):
    marker = "boundary="
    if marker not in content_type:
        raise ValueError("请上传 xlsx 文件。")
    boundary = content_type.split(marker, 1)[1].split(";", 1)[0].strip().strip('"')
    separator = ("--" + boundary).encode()
    for part in body.split(separator):
        if b"Content-Disposition" not in part or b'filename="' not in part:
            continue
        header, _, content = part.partition(b"\r\n\r\n")
        if content.endswith(b"\r\n"):
            content = content[:-2]
        disposition = header.decode("utf-8", "replace")
        filename = disposition.split('filename="', 1)[1].split('"', 1)[0]
        filename = Path(filename).name
        if not filename.lower().endswith(".xlsx"):
            raise ValueError("只接受 .xlsx 总表。")
        if not content:
            raise ValueError("文件是空的。")
        return filename, content
    raise ValueError("没有找到上传文件。")


class Handler(BaseHTTPRequestHandler):
    payload = {
        "meta": {"source": "", "as_of": "", "devices": 0, "sheets": []},
        "rows": [],
        "diff": None,
    }

    def log_message(self, fmt, *args):
        sys.stderr.write("%s - %s\n" % (self.address_string(), fmt % args))

    def send_json(self, data, status=200):
        raw = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def send_file(self, path: Path, content_type: str):
        if not path.exists():
            self.send_error(404)
            return
        raw = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/api/data":
            self.send_json(Handler.payload)
            return
        if path in ("/", "/index.html"):
            self.send_file(WEB / "index.html", "text/html; charset=utf-8")
            return
        if path == "/app.css":
            self.send_file(WEB / "app.css", "text/css; charset=utf-8")
            return
        if path == "/app.js":
            self.send_file(WEB / "app.js", "text/javascript; charset=utf-8")
            return
        self.send_error(404)

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/upload":
            self.send_error(404)
            return
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0 or length > 30 * 1024 * 1024:
            self.send_json({"error": "文件大小不合法。"}, 400)
            return
        body = self.rfile.read(length)
        try:
            filename, content = parse_upload(body, self.headers.get("Content-Type", ""))
            UPLOADS.mkdir(parents=True, exist_ok=True)
            DATA.mkdir(parents=True, exist_ok=True)
            saved = DATA / filename
            saved.write_bytes(content)
            (UPLOADS / filename).write_bytes(content)
            payload = build_from_xlsx(saved)
            save_snapshot(payload)
            Handler.payload = payload
            self.send_json(payload)
        except Exception as exc:  # noqa: BLE001 - return the parse error to the page
            self.send_json({"error": str(exc)}, 400)


def serve(port: int, open_browser: bool):
    Handler.payload = load_current()
    httpd = None
    bound = port
    for candidate in range(port, port + 20):
        try:
            httpd = ThreadingHTTPServer(("127.0.0.1", candidate), Handler)
            bound = candidate
            break
        except OSError:
            continue
    if httpd is None:
        raise SystemExit("8765-8784 端口都被占用了。")
    url = f"http://127.0.0.1:{bound}"
    meta = Handler.payload["meta"]
    print(f"付呗交易监控已打开：{url}", flush=True)
    print(f"当前总表：{meta.get('source') or '尚未导入'}，数据日 {meta.get('as_of') or '-'}", flush=True)
    if open_browser:
        webbrowser.open(url)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n已停止。")


def check():
    payload = load_current()
    rows = payload["rows"]
    counts = {}
    for row in rows:
        for flag in row["flags"]:
            counts[flag] = counts.get(flag, 0) + 1
    activated = sum(1 for row in rows if row["activated"])
    amount = round(sum(row["month_amount"] for row in rows), 2)
    tx = sum(row["month_count"] for row in rows)
    print(json.dumps({"meta": payload["meta"], "activated": activated, "amount": amount, "tx": tx, "flags": counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--check" in args:
        check()
    else:
        chosen = 8765
        browser = "--no-browser" not in args
        serve(chosen, browser)
